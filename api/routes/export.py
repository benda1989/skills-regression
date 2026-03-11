"""GET /api/export/{id} - 导出 Word 报告"""

from __future__ import annotations

import io
import logging
from typing import Any, Dict

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse

from ..auth import require_auth
from ..database import get_record

logger = logging.getLogger("regression_agent.api.export")
router = APIRouter(dependencies=[Depends(require_auth)])


@router.get(
    "/export/{record_id}",
    summary="导出分析结果为 Word 文档",
    response_class=StreamingResponse,
)
async def export_word(record_id: str) -> StreamingResponse:
    record = get_record(record_id)
    if record is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="记录不存在")

    try:
        doc_bytes = _build_word_doc(record)
    except Exception as e:
        logger.exception("Word export failed for record %s", record_id)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Word 文档生成失败：{e}",
        )

    filename = f"analysis_{record_id[:8]}.docx"
    return StreamingResponse(
        io.BytesIO(doc_bytes),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _build_word_doc(record: Dict[str, Any]) -> bytes:
    """生成 Word 文档字节流"""
    from docx import Document  # type: ignore

    doc = Document()
    doc.add_heading("回归分析报告", 0)

    result = record.get("result", {})
    ctx = result.get("analysis_context", {})
    target = ctx.get("target", "N/A")
    features = ctx.get("features", [])

    # 基本信息
    doc.add_heading("基本信息", 1)
    p = doc.add_paragraph()
    p.add_run("文件：").bold = True
    p.add_run(record.get("filename", "N/A"))
    p = doc.add_paragraph()
    p.add_run("指令：").bold = True
    p.add_run(record.get("instruction") or "（无）")
    p = doc.add_paragraph()
    p.add_run("分析时间：").bold = True
    p.add_run(record.get("created_at", "N/A"))

    # 变量设置
    doc.add_heading("变量设置", 1)
    p = doc.add_paragraph()
    p.add_run("因变量：").bold = True
    p.add_run(target)
    p = doc.add_paragraph()
    p.add_run("自变量：").bold = True
    p.add_run(", ".join(features) if features else "N/A")

    # 模型结果
    model_results = result.get("model_results", {})
    if model_results:
        doc.add_heading("模型结果", 1)
        model_type = model_results.get("model_type", "N/A")
        r2 = model_results.get("r_squared")
        adj_r2 = model_results.get("adjusted_r_squared")
        doc.add_paragraph(f"模型类型：{model_type}")
        if r2 is not None:
            doc.add_paragraph(f"R²：{r2:.4f}")
        if adj_r2 is not None:
            doc.add_paragraph(f"调整 R²：{adj_r2:.4f}")

        # 系数表
        coefficients = model_results.get("coefficients", {})
        p_values = model_results.get("p_values", {})
        if coefficients:
            doc.add_heading("回归系数", 2)
            table = doc.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            hdr[0].text = "变量"
            hdr[1].text = "系数（B）"
            hdr[2].text = "p 值"
            for feat, coef in coefficients.items():
                row = table.add_row().cells
                row[0].text = feat
                row[1].text = f"{coef:.4f}"
                pv = p_values.get(feat)
                row[2].text = f"{pv:.4f}" if pv is not None else "N/A"

    # 解读
    interpretation = result.get("interpretation", {})
    if interpretation:
        doc.add_heading("LLM 分析解读", 1)
        conclusion = interpretation.get("conclusion", {})
        if conclusion:
            summary = conclusion.get("detailed_summary", "")
            if summary:
                doc.add_paragraph(summary)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
